"""FR-021/023: carga csv/xlsx en DuckDB por espacio, joins nativos, solo lectura, timeout."""
import io

import duckdb
import pandas as pd
import pytest

from app.store import EmptyTable, QueryTimeout, SpaceExists, SpaceNotFound, SpaceStore, UnsupportedFormat

CSV = b"depto,vendedor,monto\nVentas,ana,100\nVentas,luis,50\nMarketing,jo,300\n"


def xlsx_bytes() -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame({"vendedor": ["ana", "luis", "jo"], "Zona Norte?": ["N", "S", "N"]}).to_excel(w, sheet_name="Vendedores", index=False)
        pd.DataFrame({"a": [1, 2]}).to_excel(w, sheet_name="Otra hoja", index=False)
    return buf.getvalue()


@pytest.fixture
def store(tmp_path):
    return SpaceStore(str(tmp_path), sample_rows=2)


def test_crear_espacio_y_duplicado(store):
    store.create_space("ws-1")
    assert store.exists("ws-1")
    with pytest.raises(SpaceExists):
        store.create_space("ws-1")
    with pytest.raises(ValueError):
        store.create_space("../fuera")
    with pytest.raises(SpaceNotFound):
        store.list_files("no-existe")


def test_csv_y_xlsx_join(store):
    store.create_space("ws")
    e1 = store.load_file("ws", "ventas.csv", CSV)
    assert e1["tables"][0]["name"] == "f1_ventas"
    assert e1["tables"][0]["rows"] == 3
    assert [c["name"] for c in e1["tables"][0]["columns"]] == ["depto", "vendedor", "monto"]
    assert len(e1["tables"][0]["sample"]) == 2

    e2 = store.load_file("ws", "Vendedores 2026.xlsx", xlsx_bytes())
    names = [t["name"] for t in e2["tables"]]
    assert names == ["f2_vendedores_2026_vendedores", "f2_vendedores_2026_otra_hoja"]
    cols = [c["name"] for c in e2["tables"][0]["columns"]]
    assert cols == ["vendedor", "zona_norte"]  # columnas normalizadas

    assert store.table_names("ws") == set(names) | {"f1_ventas", "t1", "t2", "t3"}
    assert e1["tables"][0]["alias"] == "t1" and [t["alias"] for t in e2["tables"]] == ["t2", "t3"]
    schema = store.schema_text("ws")
    assert "Tabla t1" in schema and "ventas.csv" in schema and "Filas de ejemplo" in schema

    columns, rows = store.query(
        "ws",
        "SELECT z.zona_norte, SUM(v.monto) AS total FROM f1_ventas v "
        "JOIN f2_vendedores_2026_vendedores z ON z.vendedor = v.vendedor GROUP BY 1 ORDER BY 1",
        timeout_s=5)
    assert columns == ["zona_norte", "total"]
    assert rows == [{"zona_norte": "N", "total": 400}, {"zona_norte": "S", "total": 50}]


def test_borrar_archivo_saca_tablas(store):
    store.create_space("ws")
    e = store.load_file("ws", "ventas.csv", CSV)
    assert store.delete_file("ws", e["file_id"]) is True
    assert store.table_names("ws") == set()
    assert store.delete_file("ws", "nope") is False


def test_formatos(store):
    store.create_space("ws")
    with pytest.raises(UnsupportedFormat):
        store.load_file("ws", "x.pdf", b"%PDF")
    with pytest.raises(EmptyTable):
        store.load_file("ws", "vacio.csv", b"")


def test_conexion_de_consulta_es_solo_lectura(store):
    store.create_space("ws")
    store.load_file("ws", "ventas.csv", CSV)
    # Aunque el validador fallara, la conexión de consulta no puede escribir ni leer archivos.
    with pytest.raises(duckdb.Error):
        store.query("ws", "CREATE TABLE hack AS SELECT 1", timeout_s=5)
    with pytest.raises(duckdb.Error):
        store.query("ws", "SELECT * FROM read_csv('/etc/passwd')", timeout_s=5)
    assert store.table_names("ws") == {"f1_ventas", "t1"}


def test_timeout_interrumpe(store):
    store.create_space("ws")
    store.load_file("ws", "ventas.csv", CSV)
    with pytest.raises(QueryTimeout):
        store.query("ws", "SELECT COUNT(*) FROM range(200000000) a, range(200000000) b", timeout_s=1)


def test_columnas_con_acentos_y_rotulo_original(store):
    store.create_space("ws")
    csv = "Número de material,Nombre Almacén,Decisión de Empleo\n1,D- Disponible,APRO\n".encode("utf-8")
    e = store.load_file("ws", "stock.csv", csv)
    cols = e["tables"][0]["columns"]
    assert [c["name"] for c in cols] == ["numero_de_material", "nombre_almacen", "decision_de_empleo"]
    assert cols[1]["label"] == "Nombre Almacén"
    schema = store.schema_text("ws")
    assert 'nombre_almacen: VARCHAR — en la planilla se llama "Nombre Almacén"' in schema
    assert "CREATE TABLE" not in schema


def test_common_keys(store):
    store.create_space("ws")
    store.load_file("ws", "ventas.csv", CSV)
    store.load_file("ws", "vendedores.xlsx", xlsx_bytes())
    assert store.common_keys("ws") == {"vendedor": ["t1", "t2"]}


def test_floats_redondeados():
    from app.store import _json_safe
    assert _json_safe(5554422.999999899) == 5554423.0
    assert _json_safe(115910370208.62003) == 115910370208.62
    assert _json_safe(float("nan")) is None


def test_encabezado_real_debajo_de_titulo(store):
    """Caso real (TABLA_DATOS.xlsx de CDCV): título en B2, nota en I4, encabezado en fila 5,
    columna A vacía. El encabezado debe detectarse y lo de arriba descartarse."""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Recogida de Datos"
    ws["B2"] = "Recogida de Datos - Proyecto"; ws["B3"] = "CÁMARA VALENCIA"; ws["I4"] = "Resaltar los más relevantes"
    for j, h in enumerate(["Año", "Fuente", "Nombre", "Formato", "Relevancia"], start=2):
        ws.cell(row=5, column=j, value=h)
    ws.append([]); 
    for r, row in enumerate([[2024, "EUROPEAN COMISSION", "DESI INDICATORS", "WEB", 90],
                             [2023, "INE", "Tipos de IA", "XSLX", 80]], start=6):
        for j, v in enumerate(row, start=2):
            ws.cell(row=r, column=j, value=v)
    buf = io.BytesIO(); wb.save(buf)
    store.create_space("ws")
    e = store.load_file("ws", "TABLA_DATOS.xlsx", buf.getvalue())
    t = e["tables"][0]
    assert [c["name"] for c in t["columns"]] == ["ano", "fuente", "nombre", "formato", "relevancia"]
    assert t["rows"] == 2
    _, rows = store.query("ws", "SELECT fuente, relevancia FROM t1 ORDER BY relevancia DESC", timeout_s=5)
    assert rows[0] == {"fuente": "EUROPEAN COMISSION", "relevancia": 90}


def test_alias_se_borra_con_el_archivo_y_se_crea_para_espacios_viejos(store):
    store.create_space("ws")
    e = store.load_file("ws", "ventas.csv", CSV)
    store.delete_file("ws", e["file_id"])
    assert store.table_names("ws") == set()
    e = store.load_file("ws", "ventas.csv", CSV)
    assert e["tables"][0]["alias"] == "t1"
    # simula un espacio cargado antes de los alias
    meta = store._read_meta("ws"); meta["files"][0]["tables"][0].pop("alias"); store._write_meta("ws", meta)
    store.ensure_aliases("ws")
    assert "t1" in store.table_names("ws")
    assert store.query("ws", "SELECT COUNT(*) AS n FROM t1", timeout_s=5)[1] == [{"n": 3}]


def test_nombre_legible_de_hoja(store):
    store.create_space("ws")
    store.load_file("ws", "libro.xlsx", xlsx_bytes())
    labels = store.table_labels("ws")
    assert labels["t1"] == 'libro.xlsx, hoja "Vendedores"' and labels["t2"] == 'libro.xlsx, hoja "Otra hoja"'
    assert 'nombre legible para mostrar: "libro.xlsx, hoja "Vendedores""' in store.schema_text("ws")


def test_diccionario_de_datos_viaja_al_esquema(store):
    store.create_space("ws")
    e = store.load_file("ws", "stock.csv", "Ce.,Alm.,UMB\nAR06,6451,UN\n".encode("utf-8"))
    n = store.set_dictionary("ws", {"t1": {"ce": "centro logístico (planta)", "umb": "unidad de medida base", "no_existe": "x"}})
    assert n == 2
    schema = store.schema_text("ws")
    assert "ce: VARCHAR — en la planilla se llama \"Ce.\" — significa: centro logístico (planta)" in schema
    assert store.list_files("ws")[0]["tables"][0]["columns"][0]["description"] == "centro logístico (planta)"
    assert store.set_dictionary("ws", {e["tables"][0]["name"]: {"ce": ""}}) == 1
    assert "description" not in store.list_files("ws")[0]["tables"][0]["columns"][0]
